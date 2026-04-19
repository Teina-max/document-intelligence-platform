"""
Import clients from 'docs/Liste client.xlsx' into Supabase entreprises table.
Uses upsert_entreprise RPC to merge with existing OCR-imported data.

Prerequisites:
  pip install openpyxl supabase

Environment:
  SUPABASE_URL — Supabase project URL
  SUPABASE_SERVICE_KEY — Service role key (bypasses RLS)

Usage:
  python scripts/import-clients-excel.py                  # dry-run (default)
  python scripts/import-clients-excel.py --execute        # actually insert
  python scripts/import-clients-excel.py --execute --verbose
"""

import os
import re
import sys
import json
import argparse
from pathlib import Path

import openpyxl


# ============================================================
# Config
# ============================================================

EXCEL_PATH = Path(__file__).parent.parent / "docs" / "Liste client.xlsx"
SHEET_NAME = "Feuil1"
HEADER_ROW = 7  # Row containing column headers
DATA_START_ROW = 8

# Map Excel countries to pays_enum values
PAYS_MAP = {
    "FR": "FR",
    "ES": "ES",
    "MAROC": "MA",
    "TUNISIE": "TN",
    "REUNION": "RE",
}

# Known bad data in conditions_paiement column (emails leaked into wrong column)
EMAIL_PATTERN = re.compile(r"^[\w.+-]+@[\w.-]+\.\w+$|^mailto:")


# ============================================================
# Helpers
# ============================================================

def clean_str(val) -> str | None:
    """Clean a cell value to a trimmed string or None."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() == "none":
        return None
    return s


def clean_email(val) -> str | None:
    """Clean an email: strip, lowercase, remove trailing dots and mailto:."""
    s = clean_str(val)
    if not s:
        return None
    s = s.lower().replace("mailto:", "").strip().rstrip(".")
    # Basic email validation
    if "@" in s and "." in s.split("@")[-1]:
        return s
    return None


def clean_phone(val) -> str | None:
    """Clean a phone number."""
    s = clean_str(val)
    if not s:
        return None
    # Remove non-digit/+ characters for validation
    digits = re.sub(r"[^\d+]", "", s)
    if len(digits) >= 8:
        return s  # Keep original formatting
    return None


def clean_conditions_paiement(val) -> str | None:
    """Clean payment terms — filter out emails that leaked into this column."""
    s = clean_str(val)
    if not s:
        return None
    if EMAIL_PATTERN.match(s):
        return None  # This is an email, not a payment term
    return s


def build_address(address: str | None, complement: str | None) -> str | None:
    """Combine address + complement into one string."""
    parts = [p for p in [clean_str(address), clean_str(complement)] if p]
    return " — ".join(parts) if parts else None


# ============================================================
# Parse Excel
# ============================================================

def parse_excel(path: Path) -> list[dict]:
    """Parse the Excel file and return a list of client dicts."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME]

    clients = []
    seen_keys = set()  # Track (nom, code_postal) to deduplicate

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        def cell(col: int):
            return ws.cell(row_idx, col).value

        numero_do = clean_str(cell(1))
        nom = clean_str(cell(2))

        if not nom:
            continue  # Skip empty rows

        pays_raw = clean_str(cell(3)) or "FR"
        pays = PAYS_MAP.get(pays_raw.upper(), "FR")

        code_postal = clean_str(cell(14))  # Billing CP

        # Deduplicate on (nom, code_postal)
        dedup_key = (nom.upper(), code_postal)
        if dedup_key in seen_keys:
            continue
        seen_keys.add(dedup_key)

        client = {
            "p_nom": nom,
            "p_numero_client": numero_do,
            "p_pays": pays,
            "p_code_postal": code_postal,
            "p_tva_intracommunautaire": clean_str(cell(4)),
            "p_siret": clean_str(cell(5)),
            "p_contact_nom": clean_str(cell(6)),
            "p_contact_telephone": clean_phone(cell(7)),
            "p_contact_email": clean_email(cell(8)),
            "p_contact_email_secondaire": clean_email(cell(9)),
            "p_email_facturation": clean_email(cell(10)),
            "p_conditions_paiement": clean_conditions_paiement(cell(11)),
            "p_adresse": build_address(cell(12), cell(13)),
            "p_ville": clean_str(cell(15)),
            "p_departement": clean_str(cell(16)),
            "p_region": clean_str(cell(17)),
            # Delivery contact
            "p_contact_livraison_nom": clean_str(cell(19)),
            "p_contact_livraison_telephone": clean_phone(cell(20)),
            "p_contact_livraison_email": clean_email(cell(21)),
            "p_adresse_livraison": build_address(cell(22), cell(23)),
            "p_code_postal_livraison": clean_str(cell(24)),
            "p_ville_livraison": clean_str(cell(25)),
            # Notes (col 28, very rare)
            "p_notes": clean_str(cell(28)),
            "p_source": "excel",
        }

        # Remove None values (let SQL defaults handle them)
        client = {k: v for k, v in client.items() if v is not None}

        clients.append(client)

    return clients


# ============================================================
# Import to Supabase
# ============================================================

def import_to_supabase(clients: list[dict], verbose: bool = False):
    """Import clients to Supabase using upsert_entreprise RPC."""
    from supabase import create_client

    url = os.environ.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_KEY")

    if not url or not key:
        print("ERROR: Set SUPABASE_URL and SUPABASE_SERVICE_KEY environment variables")
        sys.exit(1)

    sb = create_client(url, key)

    success = 0
    errors = []

    for i, client in enumerate(clients):
        try:
            result = sb.rpc("upsert_entreprise", client).execute()
            success += 1
            if verbose:
                print(f"  [{i+1}/{len(clients)}] OK: {client['p_nom']}")
        except Exception as e:
            errors.append({"client": client["p_nom"], "error": str(e)})
            print(f"  [{i+1}/{len(clients)}] ERROR: {client['p_nom']} — {e}")

    print(f"\n=== RESULT ===")
    print(f"  Success: {success}/{len(clients)}")
    if errors:
        print(f"  Errors: {len(errors)}")
        for err in errors:
            print(f"    - {err['client']}: {err['error']}")


# ============================================================
# Main
# ============================================================

def main():
    parser = argparse.ArgumentParser(description="Import ThermoPack client list to Supabase")
    parser.add_argument("--execute", action="store_true", help="Actually insert (default: dry-run)")
    parser.add_argument("--verbose", action="store_true", help="Print each client as it's processed")
    parser.add_argument("--output-json", type=str, help="Export parsed data to JSON file")
    args = parser.parse_args()

    print(f"Parsing {EXCEL_PATH}...")
    clients = parse_excel(EXCEL_PATH)

    # Stats
    pays_count = {}
    email_count = 0
    for c in clients:
        p = c.get("p_pays", "FR")
        pays_count[p] = pays_count.get(p, 0) + 1
        if "p_contact_email" in c:
            email_count += 1

    print(f"\n=== PARSED ===")
    print(f"  Total clients: {len(clients)}")
    print(f"  With email: {email_count}")
    print(f"  By country: {pays_count}")

    # Data quality warnings
    print(f"\n=== DATA QUALITY ===")
    no_email = [c["p_nom"] for c in clients if "p_contact_email" not in c]
    if no_email:
        print(f"  No email ({len(no_email)}): {', '.join(no_email[:10])}{'...' if len(no_email) > 10 else ''}")

    no_cp = [c["p_nom"] for c in clients if "p_code_postal" not in c]
    if no_cp:
        print(f"  No postal code ({len(no_cp)}): {', '.join(no_cp[:10])}{'...' if len(no_cp) > 10 else ''}")

    if args.output_json:
        with open(args.output_json, "w") as f:
            json.dump(clients, f, ensure_ascii=False, indent=2)
        print(f"\n  Exported to {args.output_json}")

    if args.execute:
        print(f"\n=== IMPORTING TO SUPABASE ===")
        import_to_supabase(clients, verbose=args.verbose)
    else:
        print(f"\n  DRY RUN — use --execute to actually import")
        if args.verbose:
            for c in clients[:5]:
                print(f"  Sample: {json.dumps(c, ensure_ascii=False)}")


if __name__ == "__main__":
    main()
